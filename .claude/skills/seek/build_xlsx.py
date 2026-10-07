"""Render a Seek JSON file into a filterable Excel workbook.

Usage: python3 build_xlsx.py seek-output/seek_<from>_<to>.json [out.xlsx]
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NO_OWNER_FILL = PatternFill("solid", fgColor="FFE699")
CATEGORY_FILLS = {
    "Work": PatternFill("solid", fgColor="DDEBF7"),
    "Personal": PatternFill("solid", fgColor="E2EFDA"),
    "Mixed": PatternFill("solid", fgColor="FCE4D6"),
}
WRAP = Alignment(wrap_text=True, vertical="top")
NO_OWNER = "No clear owner"


def write_sheet(ws, headers, widths, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def fmt_item(item):
    owner = item.get("owner") or NO_OWNER
    return f"• {item['action']} — {owner}"


def build(data, out_path):
    meetings = data["meetings"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Meetings"
    write_sheet(
        ws,
        ["Date", "Time", "Topic", "Category", "Tags", "Summary / Takeaways",
         "Action Items — Owner", "# Items", "# No Owner", "Pocket Title", "Recording ID"],
        [11, 7, 34, 10, 20, 60, 70, 8, 10, 32, 38],
        [
            [
                m["date"], m.get("time", ""), m["topic"], m["category"],
                "; ".join(m.get("tags", [])), m["summary"],
                "\n".join(fmt_item(i) for i in m["action_items"]) or "—",
                len(m["action_items"]),
                sum(1 for i in m["action_items"] if (i.get("owner") or NO_OWNER) == NO_OWNER),
                m["title"], m["recording_id"],
            ]
            for m in meetings
        ],
    )
    for row in ws.iter_rows(min_row=2):
        fill = CATEGORY_FILLS.get(row[3].value)
        if fill:
            row[3].fill = fill
        if row[8].value:
            row[8].fill = NO_OWNER_FILL

    ws = wb.create_sheet("Action Items")
    write_sheet(
        ws,
        ["Date", "Topic", "Category", "Tags", "Action Item", "Owner", "Owner Basis",
         "Due", "Priority", "Status", "Type", "Source", "Recording ID"],
        [11, 30, 10, 18, 55, 18, 34, 11, 9, 10, 15, 9, 38],
        [
            [
                m["date"], m["topic"], m["category"], "; ".join(m.get("tags", [])),
                i["action"], i.get("owner") or NO_OWNER, i.get("owner_basis", ""),
                i.get("due") or "", (i.get("priority") or "").lower(), i.get("status", ""),
                i.get("type", ""), i.get("source", ""), m["recording_id"],
            ]
            for m in meetings
            for i in m["action_items"]
        ],
    )
    for row in ws.iter_rows(min_row=2):
        fill = CATEGORY_FILLS.get(row[2].value)
        if fill:
            row[2].fill = fill
        if row[5].value == NO_OWNER:
            row[5].fill = NO_OWNER_FILL
            row[5].font = Font(bold=True)

    ws = wb.create_sheet("About")
    items = [i for m in meetings for i in m["action_items"]]
    rng = data["range"]
    lines = [
        ("Range", f"{rng['from']} to {rng['to']} ({rng.get('tz', '')})"),
        ("Generated", data.get("generated", "")),
        ("Filters", json.dumps(data.get("filters", {}))),
        ("Recordings kept", len(meetings)),
        ("  Work / Personal / Mixed", " / ".join(
            str(sum(1 for m in meetings if m["category"] == c)) for c in ("Work", "Personal", "Mixed"))),
        ("Action items", len(items)),
        ("  With no clear owner", sum(1 for i in items if (i.get("owner") or NO_OWNER) == NO_OWNER)),
        ("Recordings skipped", len(data.get("skipped", []))),
    ]
    lines += [(f"  Skipped: {s['title']}", s["reason"]) for s in data.get("skipped", [])]
    lines += [("Note", n) for n in data.get("notes", [])]
    for label, value in lines:
        ws.append([label, value])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        for cell in row:
            cell.alignment = WRAP

    wb.save(out_path)
    return len(meetings), len(items)


def main():
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".xlsx")
    data = json.loads(src.read_text())
    n_meetings, n_items = build(data, out)
    print(f"Wrote {out} ({n_meetings} recordings, {n_items} action items)")


if __name__ == "__main__":
    main()
